"""Assemble extraction/<nom_du_pdf>/pNNN.json en promotions_extraites.xlsx.

Seuls des produits alimentaires figurent dans les JSON. categorie_pyramide est
dérivée du code de sous-catégorie ; les rabais sont calculés ici.
"""
import glob, json, os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
SOUSCATS = {"1.1", "1.2", "1.3", "1.4", "2.1", "2.2", "3.1", "3.2", "3.3", "3.4", "3.5",
            "4.1", "4.2", "4.3", "4.4", "4.5", "5.1", "5.2", "5.3", "5.4", "5.5", "5.6",
            "6.1", "6.2", "6.3", "6.4", "6.5", "7.1", "7.2", "7.3", "7.4", "7.5"}
COLS = ["fichier_source", "page", "produit", "categorie_pyramide", "souscategorie_pyramide",
        "type_poisson", "provenance", "label_1", "label_2", "label_3", "type_offre_1",
        "type_offre_2", "type_offre_3", "prix_initial", "prix_final", "rabais_montant",
        "rabais_pourcentage"]


def spread(values, prefix):
    values = list(values or [])
    if len(values) > 3:
        raise ValueError(f"plus de 3 {prefix}: {values}")
    return {f"{prefix}_{i}": values[i - 1] if i <= len(values) else None for i in (1, 2, 3)}


rows = []
for folder in sorted(glob.glob(os.path.join(ROOT, "extraction", "*", ""))):
    source = os.path.basename(os.path.dirname(folder)) + ".pdf"
    for path in sorted(glob.glob(os.path.join(folder, "p*.json"))):
        page = int(re.search(r"p(\d+)\.json$", path).group(1))
        for r in json.load(open(path, encoding="utf-8")):
            code = r["souscategorie_pyramide"]
            assert code in SOUSCATS, (path, code)
            assert r["type_poisson"] is None or code == "5.3", (path, r["produit"])
            pi, pf = r["prix_initial"], r["prix_final"]
            row = {"fichier_source": source, "page": page, "produit": r["produit"],
                   "categorie_pyramide": int(code.split(".")[0]),
                   "souscategorie_pyramide": code, "type_poisson": r["type_poisson"],
                   "provenance": r["provenance"], "prix_initial": pi, "prix_final": pf,
                   "rabais_montant": round(pi - pf, 2) if pi is not None and pf is not None else None,
                   "rabais_pourcentage": round((pi - pf) / pi * 100, 1) if pi is not None and pf is not None else None}
            row.update(spread(r["labels"], "label"))
            row.update(spread(r["type_offre"], "type_offre"))
            rows.append(row)

rows.sort(key=lambda r: (r["fichier_source"], r["page"]))  # tri stable

wb = Workbook()
ws = wb.active
ws.title = "Promotions"
ws.append(COLS)
for r in rows:
    ws.append([r[c] for c in COLS])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="E65C00")
    c.alignment = Alignment(wrap_text=True, vertical="center")
for i, w in enumerate([36, 6, 70, 10, 12, 11, 26, 20, 20, 16, 24, 26, 26, 12, 12, 12, 12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row[13:16]:
        c.number_format = "0.00"
    row[16].number_format = "0.0"
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
wb.save(os.path.join(ROOT, "promotions_extraites.xlsx"))
print(len(rows), "lignes")
